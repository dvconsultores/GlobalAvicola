"""GA-CUTOVER-S3-02 · `Meta.cutover_datetime` se representa en la zona de NEGOCIO.

La persistencia sigue en UTC (columna `DateTime(timezone=True)`); lo único que cambia es la
representación en la hoja `Meta` del XLSX: el corte se muestra en `America/Caracas`
(UTC-4, offset fijo sin DST) con offset explícito y **el mismo instante**.

El parser ignora `Meta` (SPEC 011): esta hoja es display para el usuario, nunca autoridad.
"""
from __future__ import annotations

import io
import os
import sys
import time
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from cutover_v2_fixtures import preparar  # noqa: E402

# (instante de entrada en UTC, representación local esperada en America/Caracas)
VECTORES = [
    ("2026-10-28T12:00:00+00:00", "2026-10-28T08:00:00-04:00"),  # mañana
    ("2026-10-28T22:00:00+00:00", "2026-10-28T18:00:00-04:00"),  # tarde (el caso de Yolinel)
    ("2026-10-29T03:30:00+00:00", "2026-10-28T23:30:00-04:00"),  # 23:30 local, día UTC siguiente
    ("2026-10-29T04:30:00+00:00", "2026-10-29T00:30:00-04:00"),  # 00:30 local
    ("2026-11-01T00:00:00+00:00", "2026-10-31T20:00:00-04:00"),  # fin de mes
    ("2027-01-01T02:00:00+00:00", "2026-12-31T22:00:00-04:00"),  # fin de año
]


def _meta(contenido_bytes: bytes) -> dict:
    wb = load_workbook(io.BytesIO(contenido_bytes))
    return {str(k): v for k, v in wb["Meta"].iter_rows(values_only=True) if k is not None}


async def _meta_corte(http_client, auth_headers, cutover_utc: str) -> str:
    r = await http_client.get("/api/v1/cutover-templates/hatchery",
                              params={"cutover_datetime": cutover_utc}, headers=auth_headers)
    assert r.status_code == 200, r.text
    return str(_meta(r.content).get("cutover_datetime", ""))


@pytest.mark.asyncio
@pytest.mark.parametrize("input_utc,esperado_local", VECTORES)
async def test_s3_meta_muestra_corte_en_hora_de_negocio(auth_headers, http_client, seeded_ids,
                                                        input_utc, esperado_local):
    """Meta muestra el corte en UTC-4 con offset explícito y el MISMO instante."""
    await preparar(seeded_ids)
    valor = await _meta_corte(http_client, auth_headers, input_utc)
    assert valor.startswith(esperado_local), f"{input_utc} ⇒ {valor!r}, esperado {esperado_local!r}"
    assert datetime.fromisoformat(valor) == datetime.fromisoformat(input_utc), "el instante no debe cambiar"
    assert datetime.fromisoformat(valor).utcoffset() == datetime.fromisoformat(esperado_local).utcoffset()


@pytest.mark.asyncio
async def test_s3_meta_estable_ante_tz_del_proceso(auth_headers, http_client, seeded_ids):
    """La representación no depende de la TZ del proceso: sale de la zona de negocio."""
    await preparar(seeded_ids)
    previa = os.environ.get("TZ")
    try:
        os.environ["TZ"] = "America/New_York"
        time.tzset()
        valor = await _meta_corte(http_client, auth_headers, "2026-10-28T22:00:00+00:00")
        assert valor.startswith("2026-10-28T18:00:00-04:00"), valor
    finally:
        if previa is None:
            os.environ.pop("TZ", None)
        else:
            os.environ["TZ"] = previa
        time.tzset()
