"""SPEC 006 · Cutover v2 — validación de dominio (maestros, relaciones, fase, consistencia, fechas).

Negativas del mandato (§33) que requieren base de datos: maestro inválido, inactivo y de otra
empresa; galpón de otra granja; fase incompatible; lote duplicado; opening existente; orden de
fechas; consistencia sin tolerancias; plantilla de otra empresa/BU/versión.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest
from sqlalchemy import select

import app.database as database

sys.path.insert(0, str(Path(__file__).parent))
from cutover_v2_fixtures import (  # noqa: E402
    GALPON, GRANJA, INCUBADORA, PLANTA, aplicar, aprobar, corte, crear_batch, fila_engorde, fila_incubadora,
    fila_reproductoras, headers_aprobador, items_por_referencia, preparar, subir, xlsx_v2,
)
from time_reference import iso_days_ago  # noqa: E402


async def _errores(client, headers, bu: str, filas: list, empresa: int, dias: int) -> tuple[dict, list[tuple]]:
    batch_id = await crear_batch(client, headers, bu, corte(dias))
    r = await subir(client, headers, batch_id, xlsx_v2(bu, filas, company_id=empresa))
    assert r.status_code == 200, r.text
    v = (await client.get(f"/api/v1/cutover-batches/{batch_id}/validation", headers=headers)).json()
    return r.json(), sorted((e["row_number"], e["field"], e["error_code"]) for e in v["errors"])


@pytest.mark.asyncio
async def test_maestros_inexistentes_e_inactivos(auth_headers, http_client, seeded_ids):
    await preparar(seeded_ids)
    from app.masters.models import Farm

    async with database.async_session() as s:
        if (await s.execute(select(Farm).where(Farm.code == "CUT2-OFF"))).scalars().first() is None:
            s.add(Farm(company_id=seeded_ids["company_id"], code="CUT2-OFF", name="Inactiva", is_active=False))
            await s.commit()
    filas = [
        fila_reproductoras("breeder", "V-1", farm_code="NO-EXISTE"),
        fila_reproductoras("breeder", "V-2", farm_code="CUT2-OFF"),
        fila_reproductoras("breeder", "V-3", house_name="GALPON-FANTASMA"),
        fila_reproductoras("breeder", "V-4", productive_phase_code="FASE-X"),
        fila_reproductoras("breeder", "V-5", genetic_line_code="LINEA-X", breed_name="RAZA-X"),
    ]
    cuerpo, errores = await _errores(http_client, auth_headers, "breeder", filas, seeded_ids["company_id"], 60)
    assert cuerpo["invalid_rows"] == 5
    assert errores == [
        (2, "farm_code", "MASTER_NOT_FOUND"), (3, "farm_code", "MASTER_INACTIVE"), (4, "house_name", "MASTER_NOT_FOUND"),
        (5, "productive_phase_code", "MASTER_NOT_FOUND"), (6, "breed_name", "BREED_NOT_FOUND"),
        (6, "genetic_line_code", "GENETIC_LINE_NOT_FOUND"),
    ]


@pytest.mark.asyncio
async def test_maestro_de_otra_empresa_no_resuelve(auth_headers, http_client, seeded_ids):
    """S3: una granja de la empresa 2 con código válido NO se resuelve desde la empresa 1."""
    await preparar(seeded_ids)
    from app.masters.models import Farm, Hatchery

    async with database.async_session() as s:
        if (await s.execute(select(Farm).where(Farm.code == "CUT2-AJENA"))).scalars().first() is None:
            s.add(Farm(company_id=seeded_ids["company_id_2"], code="CUT2-AJENA", name="Granja ajena", is_active=True))
            s.add(Hatchery(company_id=seeded_ids["company_id_2"], code="CUT2-INC-AJENA", name="Planta ajena", is_active=True))
            await s.commit()
    _, errores = await _errores(http_client, auth_headers, "breeder",
                                [fila_reproductoras("breeder", "V-AJ", farm_code="CUT2-AJENA")], seeded_ids["company_id"], 61)
    assert errores == [(2, "farm_code", "MASTER_NOT_FOUND")]
    _, errores = await _errores(http_client, auth_headers, "hatchery",
                                [fila_incubadora("V-AJ-INC", hatchery_code="CUT2-INC-AJENA")], seeded_ids["company_id"], 61)
    assert errores == [(2, "hatchery_code", "MASTER_NOT_FOUND")]


@pytest.mark.asyncio
async def test_relaciones_padre_hijo(auth_headers, http_client, seeded_ids):
    """S6: galpón existente en OTRA granja ⇒ PARENT_MISMATCH; incubadora de otra planta ⇒ PARENT_MISMATCH."""
    await preparar(seeded_ids)
    from app.masters.models import Farm, Hatchery, House, Incubator

    async with database.async_session() as s:
        otra = (await s.execute(select(Farm).where(Farm.code == "TEST-FARM-DEST"))).scalars().first()
        if (await s.execute(select(House).where(House.farm_id == otra.id, House.name == "CUT2-GALPON-OTRA"))).scalars().first() is None:
            s.add(House(farm_id=otra.id, name="CUT2-GALPON-OTRA", capacity=1000, is_active=True))
        planta2 = (await s.execute(select(Hatchery).where(Hatchery.code == "CUT2-INC-2"))).scalars().first()
        if planta2 is None:
            planta2 = Hatchery(company_id=seeded_ids["company_id"], code="CUT2-INC-2", name="Planta 2", is_active=True)
            s.add(planta2)
            await s.flush()
            s.add(Incubator(hatchery_id=planta2.id, name="CUT2-INCUBADORA-P2", is_active=True))
        await s.commit()
    _, errores = await _errores(http_client, auth_headers, "breeder",
                                [fila_reproductoras("breeder", "V-PH", house_name="CUT2-GALPON-OTRA")], seeded_ids["company_id"], 62)
    assert errores == [(2, "house_name", "PARENT_MISMATCH")]
    _, errores = await _errores(http_client, auth_headers, "hatchery",
                                [fila_incubadora("V-PH-INC", incubator_name="CUT2-INCUBADORA-P2")], seeded_ids["company_id"], 62)
    assert errores == [(2, "incubator_name", "PARENT_MISMATCH")]


@pytest.mark.asyncio
async def test_fase_incompatible_y_metrica_no_aplicable_en_cria(auth_headers, http_client, seeded_ids):
    await preparar(seeded_ids)
    filas = [
        fila_reproductoras("breeder", "V-F1", fase="ENGORDE"),
        fila_reproductoras("breeder", "V-F2", fase="CRIA", real_start_date=iso_days_ago(70)),  # huevos KNOWN en cría
        fila_engorde("V-F3", productive_phase_code="PROD"),
    ]
    _, errores = await _errores(http_client, auth_headers, "breeder", filas[:2], seeded_ids["company_id"], 63)
    assert errores == [(2, "productive_phase_code", "PHASE_INCOMPATIBLE"),
                       (3, "accumulated_egg_production", "METRIC_NOT_APPLICABLE_IN_PHASE"),
                       (3, "accumulated_eggs_to_hatchery", "METRIC_NOT_APPLICABLE_IN_PHASE")]
    _, errores = await _errores(http_client, auth_headers, "broiler", filas[2:], seeded_ids["company_id"], 63)
    assert errores == [(2, "productive_phase_code", "PHASE_INCOMPATIBLE")]


@pytest.mark.asyncio
async def test_consistencia_sin_tolerancias(auth_headers, http_client, seeded_ids):
    await preparar(seeded_ids)
    filas = [
        # alojadas − mortalidad − descartes ≠ vivas (machos): 1100 − 50 − 0 = 1050 ≠ 1000
        fila_reproductoras("breeder", "V-C1", birds_placed_males=1100, birds_placed_females=9450),
        # con una pieza UNKNOWN la ecuación no se evalúa (no se inventa certeza)
        fila_reproductoras("breeder", "V-C2", birds_placed_males=1100, birds_placed_females=None,
                           historical_culls_males=None),
        # huevos a incubadora > producidos
        fila_reproductoras("breeder", "V-C3", accumulated_egg_production=1000, accumulated_eggs_to_hatchery=1001),
    ]
    _, errores = await _errores(http_client, auth_headers, "breeder", filas, seeded_ids["company_id"], 64)
    assert errores == [(2, "live_males", "CONSISTENCY_MISMATCH"), (4, "accumulated_eggs_to_hatchery", "CONSISTENCY_MISMATCH")]

    filas = [fila_engorde("V-C4", accumulated_birds_received=10499),  # 10499 − 300 − 200 − 0 = 9999 ≠ 10000
             fila_engorde("V-C5", partial_outbound_birds="N/A", accumulated_birds_received=1)]  # N/A ⇒ no se evalúa
    _, errores = await _errores(http_client, auth_headers, "broiler", filas, seeded_ids["company_id"], 64)
    assert errores == [(2, "live_birds", "CONSISTENCY_MISMATCH")]

    filas = [fila_incubadora("V-C6", eggs_set=10001),
             fila_incubadora("V-C7", eggs_transferred=9001),
             fila_incubadora("V-C8", eggs_transferred=8000, eggs_discarded=1001),
             fila_incubadora("V-C9", eggs_transferred=5000, chicks_hatched=5001),
             fila_incubadora("V-C10", eggs_transferred=500, chicks_hatched=100, chicks_available=101)]
    _, errores = await _errores(http_client, auth_headers, "hatchery", filas, seeded_ids["company_id"], 64)
    assert errores == [(2, "eggs_set", "CONSISTENCY_MISMATCH"), (3, "eggs_transferred", "CONSISTENCY_MISMATCH"),
                       (4, "eggs_discarded", "CONSISTENCY_MISMATCH"), (5, "chicks_hatched", "CONSISTENCY_MISMATCH"),
                       (6, "chicks_available", "CONSISTENCY_MISMATCH")]


@pytest.mark.asyncio
async def test_orden_de_fechas(auth_headers, http_client, seeded_ids):
    await preparar(seeded_ids)
    filas = [fila_reproductoras("breeder", "V-D1", real_start_date=iso_days_ago(10)),            # inicio > corte (65)
             fila_reproductoras("breeder", "V-D2", phase_start_date=iso_days_ago(250)),           # fase antes del lote (200)
             fila_reproductoras("breeder", "V-D3", phase_start_date=iso_days_ago(20))]            # fase después del corte
    _, errores = await _errores(http_client, auth_headers, "breeder", filas, seeded_ids["company_id"], 65)
    assert errores == [(2, "real_start_date", "DATE_ORDER_INVALID"), (3, "phase_start_date", "DATE_ORDER_INVALID"),
                       (4, "phase_start_date", "DATE_ORDER_INVALID")]
    filas = [fila_incubadora("V-D4", reception_date=iso_days_ago(10), set_date=None, expected_hatch_date=None),
             fila_incubadora("V-D5", set_date=iso_days_ago(125)),                                  # carga antes de recepción (120)
             fila_incubadora("V-D6", set_date=iso_days_ago(118), expected_hatch_date=iso_days_ago(119))]
    _, errores = await _errores(http_client, auth_headers, "hatchery", filas, seeded_ids["company_id"], 65)
    assert errores == [(2, "reception_date", "DATE_ORDER_INVALID"), (3, "set_date", "DATE_ORDER_INVALID"),
                       (4, "expected_hatch_date", "DATE_ORDER_INVALID")]


@pytest.mark.asyncio
async def test_lote_duplicado_existente_y_opening_existente(auth_headers, http_client, seeded_ids):
    await preparar(seeded_ids)
    empresa = seeded_ids["company_id"]
    # Un lote aplicado previamente: su código legado ya no puede volver a cargarse ni reutilizarse.
    b = await crear_batch(http_client, auth_headers, "breeder", corte(66))
    assert (await subir(http_client, auth_headers, b, xlsx_v2("breeder", [fila_reproductoras("breeder", "V-PREV")], company_id=empresa))).status_code == 200
    await aprobar(http_client, auth_headers, headers_aprobador(seeded_ids), b)
    assert (await aplicar(http_client, auth_headers, b)).status_code == 200
    lote_prev = (await items_por_referencia(http_client, auth_headers, b))["V-PREV"]

    from app.masters.models import BirdTypeEnum, Lot, LotStatus

    async with database.async_session() as s:
        if (await s.execute(select(Lot).where(Lot.lot_code == "CUT2-NATIVO-SIN-OP"))).scalars().first() is None:
            s.add(Lot(company_id=empresa, lot_code="CUT2-NATIVO-SIN-OP", origin="NATIVE", bird_type=BirdTypeEnum.BREEDER,
                      status=LotStatus.ACTIVE, activation_type="normal"))
            s.add(Lot(company_id=empresa, lot_code="CUT2-NATIVO-ENGORDE", origin="NATIVE", bird_type=BirdTypeEnum.BROILER,
                      status=LotStatus.ACTIVE, activation_type="normal"))
            await s.commit()
    filas = [
        fila_reproductoras("breeder", "V-DUP"), fila_reproductoras("breeder", "V-DUP"),                 # duplicado en el batch
        fila_reproductoras("breeder", "V-PREV"),                                                        # ya existe sin reuso explícito
        fila_reproductoras("breeder", "V-REUSE-OP", existing_lot_code="V-PREV"),                        # reuso con opening
        fila_reproductoras("breeder", "V-REUSE-BU", existing_lot_code="CUT2-NATIVO-ENGORDE"),           # reuso de otra BU
        fila_reproductoras("breeder", "V-REUSE-OK", existing_lot_code="CUT2-NATIVO-SIN-OP"),            # reuso válido
    ]
    cuerpo, errores = await _errores(http_client, auth_headers, "breeder", filas, empresa, 67)
    assert errores == [(3, "legacy_lot_code", "LOT_DUPLICATE"), (4, "legacy_lot_code", "LOT_EXISTS_REUSE_REJECTED"),
                       (5, "existing_lot_code", "OPENING_ALREADY_EXISTS"), (6, "existing_lot_code", "PARENT_MISMATCH")]
    assert (cuerpo["valid_rows"], cuerpo["invalid_rows"]) == (2, 4)
    assert lote_prev is not None


@pytest.mark.asyncio
async def test_lote_fuente_de_incubadora(auth_headers, http_client, seeded_ids):
    await preparar(seeded_ids)
    empresa = seeded_ids["company_id"]
    from app.masters.models import BirdTypeEnum, Lot, LotStatus

    async with database.async_session() as s:
        if (await s.execute(select(Lot).where(Lot.lot_code == "CUT2-FUENTE-GP"))).scalars().first() is None:
            s.add(Lot(company_id=empresa, lot_code="CUT2-FUENTE-GP", origin="NATIVE", bird_type=BirdTypeEnum.GRANDPARENT,
                      status=LotStatus.ACTIVE, activation_type="normal", legacy_lot_code="LEG-GP-1"))
            await s.commit()
    filas = [fila_incubadora("V-S1", source_lot_code="CUT2-FUENTE-GP", generation="grandparent"),
             fila_incubadora("V-S2", source_lot_code="LEG-GP-1", generation="grandparent"),   # por código legado
             fila_incubadora("V-S3", source_lot_code="CUT2-FUENTE-GP", generation="breeder"),  # generación ≠ tipo del lote
             fila_incubadora("V-S4", source_lot_code="NO-EXISTE")]
    cuerpo, errores = await _errores(http_client, auth_headers, "hatchery", filas, empresa, 68)
    assert errores == [(4, "source_lot_code", "PARENT_MISMATCH"), (5, "source_lot_code", "MASTER_NOT_FOUND")]
    assert cuerpo["valid_rows"] == 2


@pytest.mark.asyncio
async def test_campo_obligatorio_de_incubadora(auth_headers, http_client, seeded_ids):
    """S7: `eggs_set`/`eggs_received`/`hatchery_code`/`reception_date` son obligatorios."""
    await preparar(seeded_ids)
    filas = [fila_incubadora("V-H1", eggs_set=None), fila_incubadora("V-H2", hatchery_code=None),
             fila_incubadora("V-H3", eggs_received="N/A"), fila_incubadora("V-H4", reception_date=None)]
    _, errores = await _errores(http_client, auth_headers, "hatchery", filas, seeded_ids["company_id"], 69)
    assert errores == [(2, "eggs_set", "REQUIRED_FIELD_MISSING"), (3, "hatchery_code", "REQUIRED_FIELD_MISSING"),
                       (4, "eggs_received", "NOT_APPLICABLE_ON_REQUIRED"), (5, "reception_date", "REQUIRED_FIELD_MISSING")]


@pytest.mark.asyncio
async def test_plantilla_de_otra_empresa_otra_bu_y_otra_version(auth_headers, http_client, seeded_ids):
    await preparar(seeded_ids)
    empresa = seeded_ids["company_id"]
    b = await crear_batch(http_client, auth_headers, "breeder", corte(70))
    r = await subir(http_client, auth_headers, b, xlsx_v2("breeder", [fila_reproductoras("breeder", "V-E1")], company_id=empresa + 1000))
    assert r.status_code == 422 and "COMPANY_MISMATCH" in r.text
    r = await subir(http_client, auth_headers, b, xlsx_v2("hatchery", [fila_incubadora("V-E2")], company_id=empresa))
    assert r.status_code == 422 and "BUSINESS_UNIT_MISMATCH" in r.text
    r = await subir(http_client, auth_headers, b, xlsx_v2("breeder", [fila_reproductoras("breeder", "V-E3")], company_id=empresa,
                                                        template_version="v1"))
    assert r.status_code == 422 and "TEMPLATE_VERSION_UNSUPPORTED" in r.text and "retirada" in r.text
    r = await subir(http_client, auth_headers, b, b"x" * (5 * 1024 * 1024 + 1), "grande.csv")
    assert r.status_code == 422 and "FILE_TOO_LARGE" in r.text
    r = await subir(http_client, auth_headers, b, b"hola", "corte.txt")
    assert r.status_code == 422 and "UNSUPPORTED_EXTENSION" in r.text
    # Nada quedó en staging tras los rechazos.
    v = (await http_client.get(f"/api/v1/cutover-batches/{b}/validation", headers=auth_headers)).json()
    assert v["total_rows"] == 0 and v["status"] == "draft"


@pytest.mark.asyncio
async def test_catalogos_de_la_plantilla_solo_empresa_y_bu(auth_headers, http_client, seeded_ids):
    """La hoja Catalogos trae los maestros activos de la empresa efectiva para la BU; nunca ajenos ni inactivos."""
    import io

    from openpyxl import load_workbook

    await preparar(seeded_ids)
    from app.masters.models import Farm

    async with database.async_session() as s:
        if (await s.execute(select(Farm).where(Farm.code == "CUT2-AJENA-CAT"))).scalars().first() is None:
            s.add(Farm(company_id=seeded_ids["company_id_2"], code="CUT2-AJENA-CAT", name="Ajena", is_active=True))
            s.add(Farm(company_id=seeded_ids["company_id"], code="CUT2-OFF-CAT", name="Inactiva", is_active=False))
            await s.commit()
    r = await http_client.get("/api/v1/cutover-templates/breeder", headers=auth_headers)
    filas = list(load_workbook(io.BytesIO(r.content))["Catalogos"].iter_rows(values_only=True))[1:]
    codigos = {(f[0], f[1]) for f in filas}
    assert ("farm", GRANJA) in codigos and ("house", GALPON) in codigos
    assert ("productive_phase", "PROD") in codigos and ("productive_phase", "CRIA") in codigos
    assert ("productive_phase", "ENGORDE") not in codigos and ("productive_phase", "INCUB") not in codigos
    assert ("farm", "CUT2-AJENA-CAT") not in codigos and ("farm", "CUT2-OFF-CAT") not in codigos
    assert not any(f[0] == "hatchery" for f in filas)
    r = await http_client.get("/api/v1/cutover-templates/hatchery", headers=auth_headers)
    filas = list(load_workbook(io.BytesIO(r.content))["Catalogos"].iter_rows(values_only=True))[1:]
    codigos = {(f[0], f[1]) for f in filas}
    assert ("hatchery", PLANTA) in codigos and ("incubator", INCUBADORA) in codigos
    assert not any(f[0] == "farm" for f in filas)
