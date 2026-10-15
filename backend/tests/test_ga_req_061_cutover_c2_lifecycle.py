"""GA-REQ-061 · T14 · C2 — ciclo crear→subir→validar (staging) (RED).

SPEC 006 (cutover v2): la fixture `_xlsx` construye ahora la plantilla **v2 de reproductoras**
(`CUTOVER_BREEDER_V2`) a partir de las mismas filas de 8 columnas del checkpoint
(`legacy, inicio, machos, hembras, mort_m, mort_f, granja, notas`), de modo que los goldens
certificados (10.000 → 9.965, lifetime 535, UNKNOWN visible) siguen vigentes sobre el
contrato nuevo. La granja vacía se resuelve a la sembrada (en v2 es obligatoria) y la fase es
`PROD`. Un noveno elemento opcional (dict) añade columnas v2 (p. ej. `existing_lot_code`).

Rojo en HEAD (`cf03253`): no existen las rutas `/cutover-batches` (404). La
Implementación añade: creación de batch (DRAFT), upload Excel→staging (jamás a
tablas operacionales, AC43/44), parser `openpyxl` con `template_version`
validada (CUT-RED-19), errores estructurados con `error_code`/`received_value`
(AC46) y re-upload idempotente por checksum (AC50, sin duplicar filas).
"""
import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent))
from time_reference import days_ago, iso_days_ago  # noqa: E402

#: Corte por defecto del ciclo C2 (computado: las fechas literales caducan solas).
CORTE_DEFECTO = days_ago(45).isoformat() + "T00:00:00+00:00"


BU_FIXTURE = "breeder"
GRANJA_SEMBRADA = "TEST-GRANJA-01"


def _xlsx(filas: list[list], *, template_version: str = "v2", business_unit: str = BU_FIXTURE) -> bytes:
    """Plantilla v2 de la BU con las filas del checkpoint (8 columnas + extras opcionales)."""
    from cutover_v2_fixtures import fila as fila_v2, xlsx_v2

    convertidas = []
    for f in filas:
        legacy, inicio, machos, hembras, mort_m, mort_f, granja, notas = f[:8]
        extras = dict(f[8]) if len(f) > 8 and isinstance(f[8], dict) else {}
        valores = dict(
            legacy_lot_code=legacy, farm_code=granja or GRANJA_SEMBRADA,
            productive_phase_code=extras.pop("productive_phase_code", "PROD"), real_start_date=inicio,
            live_males=machos, live_females=hembras,
            historical_mortality_males=mort_m, historical_mortality_females=mort_f,
            source_system="EXCEL", notes=notas or None,
        )
        valores.update(extras)
        convertidas.append(fila_v2(business_unit, **valores))
    return xlsx_v2(business_unit, convertidas, company_id=1, template_version=template_version,
                   business_unit=business_unit)


FILAS = [
    ["CUT-BR-001", iso_days_ago(60), 5000, 5000, 300, 200, None, ""],
    ["CUT-BR-002", iso_days_ago(45), 2000, 2000, None, None, None, "mortalidad histórica desconocida"],
    # SPEC 006: el inicio real debe ser ≤ corte (45 días); 50 días deja a la fila con su único error (granja).
    ["CUT-BR-003", iso_days_ago(50), 1000, 1000, 50, 40, "F-999", ""],
]


async def _crear_batch(client, headers, cutover: str = CORTE_DEFECTO, business_unit: str = BU_FIXTURE) -> int:
    r = await client.post("/api/v1/cutover-batches", headers=headers, json={
        "business_unit": business_unit,
        "cutover_datetime": cutover,
    })
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _subir(files_ok: bytes):
    return {"file": ("corte.xlsx", files_ok,
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


@pytest.mark.asyncio
async def test_c2_crear_batch_draft(auth_headers, client, seeded_ids):
    """POST /cutover-batches nace en DRAFT con actor y empresa efectiva."""
    r = await client.post("/api/v1/cutover-batches", headers=auth_headers, json={
        "business_unit": BU_FIXTURE,
        "cutover_datetime": CORTE_DEFECTO,
    })
    assert r.status_code == 201, r.text
    cuerpo = r.json()
    assert cuerpo["status"] == "draft"
    assert cuerpo["business_unit"] == BU_FIXTURE
    assert cuerpo["company_id"] == seeded_ids["company_id"]
    assert cuerpo["total_rows"] == 0


async def _fases_v2() -> None:
    """La validación v2 resuelve `productive_phase_code` contra el catálogo: las fases deben existir."""
    from cutover_v2_fixtures import asegurar_fases

    await asegurar_fases()


@pytest.mark.asyncio
async def test_c2_upload_validar_preview(auth_headers, http_client, seeded_ids):
    """Upload → staging + validación sin efectos: 3 leídas / 2 válidas / 1 con error estructurado."""
    await _fases_v2()
    batch_id = await _crear_batch(http_client, auth_headers)
    r = await http_client.post(f"/api/v1/cutover-batches/{batch_id}/upload",
                               headers=auth_headers, files=_subir(_xlsx(FILAS)))
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert (cuerpo["total_rows"], cuerpo["valid_rows"], cuerpo["invalid_rows"]) == (3, 2, 1), cuerpo
    assert cuerpo["status"] == "validating"

    # Preview (validación) con error estructurado: fila 4 (encabezado=1) · farm_code F-999 no existe.
    v = await http_client.get(f"/api/v1/cutover-batches/{batch_id}/validation", headers=auth_headers)
    assert v.status_code == 200, v.text
    errores = v.json()["errors"]
    assert len(errores) == 1
    err = errores[0]
    assert err["row_number"] == 4
    assert err["error_code"] == "MASTER_NOT_FOUND"
    assert err["field"] == "farm_code"
    assert err["received_value"] == "F-999"

    # Nada fué aplicado: sin lotes creados, sin errores de staging duplicado.
    listado = await http_client.get(
        f"/api/v1/cutover-batches/{batch_id}/items", headers=auth_headers)
    assert listado.status_code == 200, listado.text
    items = listado.json()["items"]
    assert len(items) == 3
    assert {it["validation_status"] for it in items} == {"valid", "invalid"}


@pytest.mark.asyncio
async def test_c2_template_version_no_soportada(auth_headers, http_client):
    """CUT-RED-19: versión de plantilla no soportada ⇒ error determinista, 0 aplicados."""
    batch_id = await _crear_batch(http_client, auth_headers)
    r = await http_client.post(f"/api/v1/cutover-batches/{batch_id}/upload",
                               headers=auth_headers, files=_subir(_xlsx(FILAS, template_version="v99")))
    assert r.status_code in (400, 422), r.text
    assert "TEMPLATE_VERSION_UNSUPPORTED" in r.text
    # El batch sigue sin filas de staging y en DRAFT.
    v = await http_client.get(f"/api/v1/cutover-batches/{batch_id}/validation", headers=auth_headers)
    assert v.json()["total_rows"] == 0


@pytest.mark.asyncio
async def test_c2_mismo_checksum_no_duplica(auth_headers, http_client):
    """AC50/CUT-RED-08 (misma carga): re-subir el MISMO archivo no duplica filas.

    Corte propio: el pre-chequeo de duplicado es por (empresa, BU, checksum,
    corte) — compartir corte con otro test haría que este viera un 409 ajeno.
    """
    await _fases_v2()
    batch_id = await _crear_batch(http_client, auth_headers, cutover=days_ago(43).isoformat() + "T00:00:00+00:00")
    contenido = _xlsx(FILAS)
    r1 = await http_client.post(f"/api/v1/cutover-batches/{batch_id}/upload",
                                headers=auth_headers, files=_subir(contenido))
    assert r1.status_code == 200, r1.text
    r2 = await http_client.post(f"/api/v1/cutover-batches/{batch_id}/upload",
                                headers=auth_headers, files=_subir(contenido))
    assert r2.status_code == 200, r2.text
    assert r2.json()["total_rows"] == 3  # no 6
    listado = await http_client.get(f"/api/v1/cutover-batches/{batch_id}/items", headers=auth_headers)
    assert len(listado.json()["items"]) == 3


@pytest.mark.asyncio
async def test_c2_plantilla_descargable_y_cargable(auth_headers, http_client):
    """`GET /cutover-templates/{bu}`: plantilla versionada descargable; el propio archivo re-cargado se valida.

    SPEC 006 (Q6): la versión vigente es **v2** (v1 retirada) y la plantilla trae `Catalogos`.
    Subir la plantilla vacía ya no deja un batch VALIDATED sin filas: `NO_DATA_ROWS` (RED-01).
    """
    from openpyxl import load_workbook

    for bu in ("grandparent", "breeder", "hatchery", "broiler"):
        r = await http_client.get(f"/api/v1/cutover-templates/{bu}", headers=auth_headers)
        assert r.status_code == 200, r.text
        assert "spreadsheetml" in r.headers["content-type"]
        wb = load_workbook(io.BytesIO(r.content))
        assert {"Instrucciones", "Meta", "Datos", "Catalogos"} <= set(wb.sheetnames)
        meta = {fila[0].value: fila[1].value for fila in wb["Meta"].iter_rows(min_row=2)}
        assert meta["template_version"] == "v2"
        assert meta["business_unit"] == bu

    # La plantilla descargada sin filas no se acepta como carga válida (sin efectos).
    r = await http_client.get(f"/api/v1/cutover-templates/{BU_FIXTURE}", headers=auth_headers)
    batch_id = await _crear_batch(http_client, auth_headers)
    subida = await http_client.post(f"/api/v1/cutover-batches/{batch_id}/upload",
                                    headers=auth_headers, files={"file": ("plantilla.xlsx", r.content,
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert subida.status_code == 422, subida.text
    assert "NO_DATA_ROWS" in subida.text
    v = await http_client.get(f"/api/v1/cutover-batches/{batch_id}/validation", headers=auth_headers)
    assert v.json()["total_rows"] == 0 and v.json()["status"] == "draft"

    # Unidad inexistente ⇒ 422 determinista.
    r = await http_client.get("/api/v1/cutover-templates/aves", headers=auth_headers)
    assert r.status_code == 422
    assert "BUSINESS_UNIT_INVALID" in r.text
