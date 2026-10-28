"""SPEC 006 · Cutover v2 — suite RED (mandato §34; 9 escenarios).

Rojo en HEAD (`57a421a`): una sola plantilla genérica v1 para las cuatro BU, sin CSV,
sin contratos por BU, N/A y vacío colapsados, incubadora irrepresentable, sin listado de
batches, sin informe de errores descargable, sin CUT-GL-02 y saldos de huevos/pollitos
que ignoran el opening. Cada prueba afirma el comportamiento correcto del contrato v2
(`specs/006-cutover-initial-balances-remediation/`), no un detalle de implementación.
"""
from __future__ import annotations

import io
import sys
from datetime import datetime
from pathlib import Path

import pytest
from sqlalchemy import select

import app.database as database

sys.path.insert(0, str(Path(__file__).parent))
from cutover_v2_fixtures import (  # noqa: E402
    COLUMNAS, GALPON, GRANJA, SCHEMA_IDS, TEMPLATE_VERSIONS, aplicar, aprobar, asegurar_fases, corte, crear_batch,
    csv_v2,
    fila, fila_engorde, fila_incubadora, fila_reproductoras, headers_aprobador, items_por_referencia,
    preparar, subir, xlsx_v2,
)
from time_reference import iso_days_ago, recent_event_date  # noqa: E402


def _hojas(contenido: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(contenido), data_only=True)


def _cabecera(wb) -> list[str]:
    return [c for c in next(wb["Datos"].iter_rows(values_only=True)) if c is not None]


def _meta(wb) -> dict:
    return {str(k): v for k, v in wb["Meta"].iter_rows(values_only=True) if k is not None}


# ── RED-01 · fallo en vivo reproducido ───────────────────────────────────────

@pytest.mark.asyncio
async def test_red_01_fallo_en_vivo_plantilla_vacia_errores_invisibles_y_batch_perdido(
        auth_headers, http_client, seeded_ids):
    """Lo que vio el Owner: archivo vacío «validado», errores sin mensaje, batch que desaparece.

    Contrato v2: 0 filas ⇒ `NO_DATA_ROWS` (nunca VALIDATED); cada error trae `message`;
    el informe de errores se descarga en CSV; el batch se lista y se consulta por id; la
    plantilla pedida con `batch_id` de otra BU se rechaza (la BU la fija el batch).
    """
    await preparar(seeded_ids)
    empresa = seeded_ids["company_id"]
    c = corte(41)
    batch_id = await crear_batch(http_client, auth_headers, "breeder", c)

    # (a) plantilla sin filas: no puede quedar VALIDATED.
    r = await subir(http_client, auth_headers, batch_id, xlsx_v2("breeder", [], company_id=empresa))
    assert r.status_code == 422, r.text
    assert "NO_DATA_ROWS" in r.text

    # (b) errores «de negocio» con mensaje y descargables.
    filas = [fila("breeder", legacy_lot_code="LOTE-X", farm_code="GRANJA-X", productive_phase_code="PROD",
                  real_start_date="01/08/2026", live_males="1.200", live_females=800)]
    r = await subir(http_client, auth_headers, batch_id, xlsx_v2("breeder", filas, company_id=empresa))
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "validating"
    v = (await http_client.get(f"/api/v1/cutover-batches/{batch_id}/validation", headers=auth_headers)).json()
    codigos = {e["error_code"] for e in v["errors"]}
    assert {"INVALID_DATE", "INVALID_INTEGER", "MASTER_NOT_FOUND"} <= codigos, codigos
    assert all(e["message"] for e in v["errors"])
    informe = await http_client.get(f"/api/v1/cutover-batches/{batch_id}/errors?format=csv", headers=auth_headers)
    assert informe.status_code == 200, informe.text
    assert informe.headers["content-type"].startswith("text/csv")
    texto = informe.content.decode("utf-8-sig")
    assert texto.splitlines()[0].split(";")[:5] == ["row_number", "field", "error_code", "message", "received_value"]
    assert "GRANJA-X" in texto and "INVALID_DATE" in texto

    # (c) persistencia visible: listado y detalle por id.
    listado = await http_client.get("/api/v1/cutover-batches", headers=auth_headers)
    assert listado.status_code == 200, listado.text
    assert any(b["id"] == batch_id for b in listado.json()["items"])
    detalle = await http_client.get(f"/api/v1/cutover-batches/{batch_id}", headers=auth_headers)
    assert detalle.status_code == 200 and detalle.json()["business_unit"] == "breeder"

    # (d) la plantilla sigue al batch, no a un selector.
    cruzada = await http_client.get(f"/api/v1/cutover-templates/hatchery?batch_id={batch_id}", headers=auth_headers)
    assert cruzada.status_code == 422 and "BUSINESS_UNIT_MISMATCH" in cruzada.text


# ── RED-02 · plantilla genérica insuficiente ─────────────────────────────────

@pytest.mark.asyncio
async def test_red_02_plantilla_por_bu_no_generica(auth_headers, http_client):
    """Cuatro plantillas distintas, versión v2, con `Instrucciones`, `Meta`, `Datos`, `Catalogos`."""
    cabeceras: dict[str, list[str]] = {}
    for bu in ("grandparent", "breeder", "hatchery", "broiler"):
        r = await http_client.get(f"/api/v1/cutover-templates/{bu}", headers=auth_headers)
        assert r.status_code == 200, r.text
        assert f"cutover_template_{bu}_{TEMPLATE_VERSIONS[bu]}.xlsx" in r.headers.get("content-disposition", "")
        wb = _hojas(r.content)
        assert {"Instrucciones", "Meta", "Datos", "Catalogos"} <= set(wb.sheetnames), wb.sheetnames
        meta = _meta(wb)
        assert meta["template_version"] == TEMPLATE_VERSIONS[bu]
        assert meta["schema_id"] == SCHEMA_IDS[bu]
        assert meta["business_unit"] == bu
        cabeceras[bu] = _cabecera(wb)
        assert cabeceras[bu] == list(COLUMNAS[bu]), (bu, cabeceras[bu])
    assert cabeceras["hatchery"] != cabeceras["breeder"] != cabeceras["broiler"]
    assert cabeceras["grandparent"] == cabeceras["breeder"]  # mismo contrato de reproductoras


# ── RED-03 · campos específicos de la BU ausentes ────────────────────────────

@pytest.mark.asyncio
async def test_red_03_campos_especificos_por_bu_presentes(auth_headers, http_client):
    """Los campos que el dominio necesita por BU existen en su plantilla (no en `notes`)."""
    esperados = {
        "breeder": {"productive_phase_code", "house_name", "historical_culls_males", "accumulated_feed_kg",
                    "current_avg_weight_g", "accumulated_egg_production", "accumulated_eggs_to_hatchery",
                    "birds_placed_males", "genetic_line_code"},
        "broiler": {"live_birds", "age_days", "accumulated_birds_received", "partial_outbound_birds",
                    "historical_mortality", "historical_culls"},
        "hatchery": {"hatchery_code", "incubator_name", "hatcher_name", "source_lot_code", "generation",
                     "reception_date", "set_date", "eggs_received", "eggs_set", "eggs_transferred",
                     "eggs_discarded", "chicks_hatched", "chicks_available"},
    }
    for bu, campos in esperados.items():
        r = await http_client.get(f"/api/v1/cutover-templates/{bu}", headers=auth_headers)
        assert r.status_code == 200, r.text
        cabecera = set(_cabecera(_hojas(r.content)))
        assert campos <= cabecera, (bu, sorted(campos - cabecera))
    # La hoja Instrucciones describe cada campo del contrato (una fila por campo).
    r = await http_client.get("/api/v1/cutover-templates/hatchery", headers=auth_headers)
    instrucciones = [str(f[0]) for f in _hojas(r.content)["Instrucciones"].iter_rows(values_only=True) if f and f[0]]
    for campo in esperados["hatchery"]:
        assert any(campo in linea for linea in instrucciones), campo


# ── RED-04 · CSV no soportado ────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_red_04_csv_se_descarga_y_se_carga_como_xlsx(auth_headers, http_client, seeded_ids):
    """CSV v2: plantilla descargable y carga válida; misma normalización que XLSX."""
    await preparar(seeded_ids)
    empresa = seeded_ids["company_id"]
    plantilla = await http_client.get("/api/v1/cutover-templates/broiler?format=csv", headers=auth_headers)
    assert plantilla.status_code == 200, plantilla.text
    assert plantilla.headers["content-type"].startswith("text/csv")
    cabecera = plantilla.content.decode("utf-8-sig").splitlines()[0]
    assert cabecera.split(";") == list(COLUMNAS["broiler"]) or cabecera.split(",") == list(COLUMNAS["broiler"])

    filas = [fila_engorde("CUT2-CSV-1")]
    b_csv = await crear_batch(http_client, auth_headers, "broiler", corte(42))
    r = await subir(http_client, auth_headers, b_csv, csv_v2("broiler", filas), "corte.csv")
    assert r.status_code == 200, r.text
    assert (r.json()["status"], r.json()["total_rows"], r.json()["valid_rows"]) == ("validated", 1, 1)
    assert r.json()["source_type"] == "CSV"

    b_xlsx = await crear_batch(http_client, auth_headers, "broiler", corte(43))
    r2 = await subir(http_client, auth_headers, b_xlsx, xlsx_v2("broiler", filas, company_id=empresa))
    assert r2.status_code == 200 and r2.json()["status"] == "validated", r2.text

    # Paridad: la fila normalizada es idéntica en ambos formatos.
    from app.cutover.models import CutoverStagingRow

    async with database.async_session() as s:
        n_csv = (await s.execute(select(CutoverStagingRow.normalized).where(
            CutoverStagingRow.batch_id == b_csv))).scalar_one()
        n_xlsx = (await s.execute(select(CutoverStagingRow.normalized).where(
            CutoverStagingRow.batch_id == b_xlsx))).scalar_one()
    assert n_csv == n_xlsx
    assert n_csv["live_birds"] == {"value": 10000, "status": "KNOWN"}

    # Cabecera obligatoria ausente en CSV ⇒ rechazo (S12).
    sin_columna = tuple(c for c in COLUMNAS["broiler"] if c != "live_birds")
    b_mal = await crear_batch(http_client, auth_headers, "broiler", corte(44))
    r3 = await subir(http_client, auth_headers, b_mal,
                     csv_v2("broiler", [fila_engorde("CUT2-CSV-2")[:11] + fila_engorde("CUT2-CSV-2")[12:]],
                            columnas=sin_columna), "corte.csv")
    assert r3.status_code == 422 and "HEADER_MISSING" in r3.text, r3.text


# ── RED-05 · incubadora no representable ─────────────────────────────────────

@pytest.mark.asyncio
async def test_red_05_incubadora_representada_y_aplicada(auth_headers, client, http_client, seeded_ids):
    """Una fila de incubadora produce `hatchery_openings` y saldos de huevo reales."""
    ids = await preparar(seeded_ids)
    empresa = seeded_ids["company_id"]
    c = corte(45)
    batch_id = await crear_batch(http_client, auth_headers, "hatchery", c)
    r = await subir(http_client, auth_headers, batch_id,
                    xlsx_v2("hatchery", [fila_incubadora("CUT2-INC-1")], company_id=empresa))
    assert r.status_code == 200 and r.json()["status"] == "validated", r.text
    await aprobar(http_client, auth_headers, headers_aprobador(seeded_ids), batch_id)
    r = await aplicar(http_client, auth_headers, batch_id)
    assert r.status_code == 200, r.text

    lote = (await items_por_referencia(http_client, auth_headers, batch_id))["CUT2-INC-1"]
    from app.cutover.models import HatcheryOpening
    from app.lots.models import LotPhase, OpeningBalance
    from app.masters.models import Lot
    from app.operations.validators import get_hatchery_egg_balance

    async with database.async_session() as s:
        l = (await s.execute(select(Lot).where(Lot.id == lote))).scalar_one()
        assert l.bird_type.value == "hatchery" and l.origin == "MIGRATED" and l.hatchery_purpose == "breeder"
        ho = (await s.execute(select(HatcheryOpening).where(HatcheryOpening.lot_id == lote))).scalar_one()
        assert (ho.hatchery_id, ho.incubator_id, ho.hatcher_id) == (
            ids["hatchery_id"], ids["incubator_id"], ids["hatcher_id"])
        assert (ho.eggs_received, ho.eggs_set, ho.eggs_discarded) == (10000, 9000, 100)
        assert ho.eggs_transferred == 0 and ho.eggs_transferred_status == "KNOWN"
        ob = (await s.execute(select(OpeningBalance).where(OpeningBalance.lot_id == lote))).scalar_one()
        assert ob.phase_at_activation_id == ids["INCUB"]
        assert ob.chicks_hatched_status == "KNOWN" and ob.accumulated_chicks_hatched == 0
        fase = (await s.execute(select(LotPhase).where(LotPhase.lot_id == lote))).scalar_one()
        assert fase.phase_id == ids["INCUB"] and fase.is_active
        assert await get_hatchery_egg_balance(s, lote) == 1000  # recibidos − cargados


# ── RED-06 · upload real desde la plantilla descargada ───────────────────────

@pytest.mark.asyncio
async def test_red_06_plantilla_descargada_rellenada_y_cargada(auth_headers, http_client, seeded_ids):
    """Ida y vuelta real: descargar con `batch_id`, rellenar `Datos`, subir ⇒ 0 errores + preview por BU."""
    from openpyxl import load_workbook

    await preparar(seeded_ids)
    c = corte(46)
    batch_id = await crear_batch(http_client, auth_headers, "breeder", c)
    r = await http_client.get(f"/api/v1/cutover-templates/breeder?batch_id={batch_id}", headers=auth_headers)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    meta = _meta(wb)
    # GA-CUTOVER-S3-02: Meta representa el corte en la zona de negocio (UTC-4, `-04:00`); el
    # instante sigue siendo el del batch, así que se comparan instantes y no la hora de pared.
    assert meta["template_version"] == "v2"
    assert datetime.fromisoformat(str(meta["cutover_datetime"])) == datetime.fromisoformat(corte(46))
    catalogos = list(wb["Catalogos"].iter_rows(values_only=True))
    assert any(f[0] == "farm" and f[1] == GRANJA for f in catalogos[1:]), "catálogo de granjas de la empresa"
    assert any(f[0] == "house" and f[1] == GALPON and f[4] == GRANJA for f in catalogos[1:]), "galpón con su granja"
    assert any(f[0] == "productive_phase" and f[1] == "PROD" for f in catalogos[1:])

    wb["Datos"].append(fila_reproductoras("breeder", "CUT2-DL-1", fase="PROD"))
    wb["Datos"].append(fila_reproductoras("breeder", "CUT2-DL-2", fase="CRIA", real_start_date=iso_days_ago(70),
                                          accumulated_egg_production="N/A", accumulated_eggs_to_hatchery="N/A",
                                          live_males=500, live_females=4500))
    buf = io.BytesIO()
    wb.save(buf)
    r = await subir(http_client, auth_headers, batch_id, buf.getvalue(), "cutover_template_breeder_v2.xlsx")
    assert r.status_code == 200, r.text
    assert (r.json()["status"], r.json()["total_rows"], r.json()["invalid_rows"]) == ("validated", 2, 0), r.text
    preview = await http_client.get(f"/api/v1/cutover-batches/{batch_id}/preview", headers=auth_headers)
    assert preview.status_code == 200, preview.text
    cuerpo = preview.json()
    assert cuerpo["business_unit"] == "breeder"
    assert cuerpo["summary"]["lots"] == 2
    assert cuerpo["summary"]["live_birds"] == 15000
    assert len(cuerpo["rows"]) == 2


# ── RED-07 · pérdida UNKNOWN vs N/A ──────────────────────────────────────────

@pytest.mark.asyncio
async def test_red_07_unknown_cero_y_na_se_distinguen(auth_headers, http_client, seeded_ids):
    """vacío ⇒ UNKNOWN · 0 ⇒ KNOWN 0 · N/A ⇒ NOT_APPLICABLE, en staging, opening y reconciliación.
    N/A en un campo obligatorio ⇒ error, jamás 0."""
    ids = await preparar(seeded_ids)
    empresa = seeded_ids["company_id"]
    c = corte(47)
    batch_id = await crear_batch(http_client, auth_headers, "grandparent", c)
    filas = [fila_reproductoras("grandparent", "CUT2-SEM-1", fase="CRIA", real_start_date=iso_days_ago(70),
                                historical_mortality_males=None, historical_mortality_females=None,
                                historical_culls_males=0, historical_culls_females=0,
                                accumulated_egg_production="N/A", accumulated_eggs_to_hatchery="N/A",
                                accumulated_feed_kg=None)]
    r = await subir(http_client, auth_headers, batch_id, xlsx_v2("grandparent", filas, company_id=empresa))
    assert r.status_code == 200 and r.json()["status"] == "validated", r.text

    from app.cutover.models import CutoverStagingRow

    async with database.async_session() as s:
        n = (await s.execute(select(CutoverStagingRow.normalized).where(
            CutoverStagingRow.batch_id == batch_id))).scalar_one()
    assert n["historical_mortality_males"] == {"value": None, "status": "UNKNOWN"}
    assert n["historical_culls_males"] == {"value": 0, "status": "KNOWN"}
    assert n["accumulated_egg_production"] == {"value": None, "status": "NOT_APPLICABLE"}
    assert n["accumulated_feed_kg"] == {"value": None, "status": "UNKNOWN"}

    await aprobar(http_client, auth_headers, headers_aprobador(seeded_ids), batch_id)
    assert (await aplicar(http_client, auth_headers, batch_id)).status_code == 200
    lote = (await items_por_referencia(http_client, auth_headers, batch_id))["CUT2-SEM-1"]
    from app.lots.models import OpeningBalance

    async with database.async_session() as s:
        ob = (await s.execute(select(OpeningBalance).where(OpeningBalance.lot_id == lote))).scalar_one()
    assert ob.mortality_status == "UNKNOWN"
    assert ob.culls_status == "KNOWN" and (ob.accumulated_culls_male, ob.accumulated_culls_female) == (0, 0)
    assert ob.egg_production_status == "NOT_APPLICABLE"
    assert ob.feed_status == "UNKNOWN" and ob.accumulated_feed_kg is None
    assert ob.phase_at_activation_id == ids["CRIA"]

    rec = (await http_client.get(f"/api/v1/cutover-batches/{batch_id}/reconciliation", headers=auth_headers)).json()
    f = rec["lots"][0]
    assert f["opening"]["historical_mortality"] is None and f["opening"]["mortality_status"] == "UNKNOWN"
    assert f["opening"]["egg_production_status"] == "NOT_APPLICABLE"
    assert f["lifetime"]["mortality"] is None

    # N/A donde el dato es obligatorio: error estructurado, no cero.
    b2 = await crear_batch(http_client, auth_headers, "breeder", corte(48))
    r = await subir(http_client, auth_headers, b2, xlsx_v2(
        "breeder", [fila_reproductoras("breeder", "CUT2-SEM-2", live_males="N/A")], company_id=empresa))
    assert r.status_code == 200 and r.json()["invalid_rows"] == 1, r.text
    v = (await http_client.get(f"/api/v1/cutover-batches/{b2}/validation", headers=auth_headers)).json()
    assert any(e["error_code"] == "NOT_APPLICABLE_ON_REQUIRED" and e["field"] == "live_males" for e in v["errors"])


# ── RED-08 · campo sin destino ───────────────────────────────────────────────

def test_red_08_registry_cerrado_sin_campos_sin_destino():
    """`UNMAPPED_TEMPLATE_FIELDS = 0` y `REQUIRED_DOMAIN_FIELDS_MISSING_FROM_TEMPLATE = 0`."""
    from app.cutover.contracts import CONTRACTS, campos_dominio_requeridos, lineage

    assert set(CONTRACTS) == {"grandparent", "breeder", "hatchery", "broiler"}
    for bu, contrato in CONTRACTS.items():
        assert contrato.schema_id == SCHEMA_IDS[bu]
        assert contrato.version == TEMPLATE_VERSIONS[bu]
        assert tuple(f.name for f in contrato.fields) == COLUMNAS[bu], bu
        sin_destino = [f.name for f in contrato.fields if not f.targets]
        assert sin_destino == [], (bu, sin_destino)
        faltantes = campos_dominio_requeridos(bu) - {f.name for f in contrato.fields}
        assert faltantes == set(), (bu, faltantes)
    matriz = lineage()
    assert matriz["UNMAPPED_TEMPLATE_FIELDS"] == 0
    assert matriz["REQUIRED_DOMAIN_FIELDS_MISSING_FROM_TEMPLATE"] == 0


# ── RED-09 · primera operación tras el corte ─────────────────────────────────

@pytest.mark.asyncio
async def test_red_09_primera_operacion_post_corte_consume_el_opening(auth_headers, http_client, seeded_ids):
    """Engorde: mortalidad sobre un opening mixto; incubadora: carga de incubación sobre el huevo
    recibido; CUT-GL-02: evento con fecha ≤ corte rechazado, > corte aceptado."""
    await preparar(seeded_ids)
    empresa = seeded_ids["company_id"]
    c = corte(40)

    # Engorde ────────────────────────────────────────────────────────────────
    b = await crear_batch(http_client, auth_headers, "broiler", c)
    r = await subir(http_client, auth_headers, b, xlsx_v2("broiler", [fila_engorde("CUT2-ENG-1")], company_id=empresa))
    assert r.status_code == 200 and r.json()["status"] == "validated", r.text
    await aprobar(http_client, auth_headers, headers_aprobador(seeded_ids), b)
    assert (await aplicar(http_client, auth_headers, b)).status_code == 200
    lote_eng = (await items_por_referencia(http_client, auth_headers, b))["CUT2-ENG-1"]

    from app.operations.validators import get_current_bird_balance

    async with database.async_session() as s:
        assert await get_current_bird_balance(s, lote_eng) == 10000  # mixto cuenta

    # CUT-GL-02: el día del corte y antes ⇒ rechazado con la regla; después ⇒ aceptado.
    for fecha in (iso_days_ago(40), iso_days_ago(45)):
        r = await http_client.post("/api/v1/operations", headers=auth_headers, json={
            "lot_id": lote_eng, "event_type": "mortality_recording", "event_date": fecha,
            "bird_movements": [{"sex": "mixed", "quantity": 5, "avg_weight": 900.0}]})
        assert r.status_code == 400, r.text
        assert r.json().get("rule") == "CUT-GL-02", r.text
    r = await http_client.post("/api/v1/operations", headers=auth_headers, json={
        "lot_id": lote_eng, "event_type": "mortality_recording", "event_date": recent_event_date(),
        "bird_movements": [{"sex": "mixed", "quantity": 35, "avg_weight": 900.0}]})
    assert r.status_code == 201, r.text
    rec = (await http_client.get(f"/api/v1/cutover-batches/{b}/reconciliation", headers=auth_headers)).json()
    f = next(x for x in rec["lots"] if x["lot_id"] == lote_eng)
    assert (f["opening"]["live"], f["post"]["mortality"], f["current_live"]) == (10000, 35, 9965)
    assert f["lifetime"]["mortality"] == 335  # 300 histórico KNOWN + 35

    # Incubadora ─────────────────────────────────────────────────────────────
    bh = await crear_batch(http_client, auth_headers, "hatchery", corte(39))
    r = await subir(http_client, auth_headers, bh, xlsx_v2("hatchery", [fila_incubadora("CUT2-INC-9")], company_id=empresa))
    assert r.status_code == 200 and r.json()["status"] == "validated", r.text
    await aprobar(http_client, auth_headers, headers_aprobador(seeded_ids), bh)
    assert (await aplicar(http_client, auth_headers, bh)).status_code == 200
    lote_inc = (await items_por_referencia(http_client, auth_headers, bh))["CUT2-INC-9"]
    r = await http_client.post("/api/v1/operations", headers=auth_headers, json={
        "lot_id": lote_inc, "event_type": "incubation_load", "event_date": recent_event_date(),
        "hatchery_params": [{"quantity_loaded": 1000}]})
    assert r.status_code == 201, r.text  # BR-03: hay 1.000 huevos recibidos sin cargar
    r = await http_client.post("/api/v1/operations", headers=auth_headers, json={
        "lot_id": lote_inc, "event_type": "incubation_load", "event_date": recent_event_date(),
        "hatchery_params": [{"quantity_loaded": 1}]})
    assert r.status_code == 400, r.text  # ya no queda huevo disponible: el opening se consumió
