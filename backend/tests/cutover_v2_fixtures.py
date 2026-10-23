"""SPEC 006 · Cutover v2 — fixtures compartidas (contratos por BU, XLSX/CSV, maestros).

Copia **independiente del código** de las cabeceras del contrato
(`specs/006-…/template-contract.md`): si la implementación cambia una columna, la
suite debe fallar, no heredar el cambio.

Semántica de celda: vacío = UNKNOWN · `0` = KNOWN cero · `N/A` = NOT_APPLICABLE.
"""
from __future__ import annotations

import csv
import io
import sys
from pathlib import Path
from typing import Any, Iterable, Optional

from sqlalchemy import select

import app.database as database

sys.path.insert(0, str(Path(__file__).parent))
from time_reference import days_ago, iso_days_ago  # noqa: E402

TEMPLATE_VERSION = "v2"

SCHEMA_IDS = {
    "grandparent": "CUTOVER_GRANDPARENT_V2",
    "breeder": "CUTOVER_BREEDER_V2",
    "hatchery": "CUTOVER_HATCHERY_V3",
    "broiler": "CUTOVER_BROILER_V2",
}

#: SPEC 009: la plantilla de incubadora es v3; el resto sigue en v2.
TEMPLATE_VERSIONS = {"grandparent": "v2", "breeder": "v2", "hatchery": "v3", "broiler": "v2"}

_IDENTIFICACION_AVES = (
    "legacy_lot_code", "existing_lot_code", "farm_code", "house_name", "area_code",
    "productive_phase_code", "real_start_date", "phase_start_date",
    "genetic_line_code", "breed_name", "sex",
)
_PROCEDENCIA = ("source_system", "source_reference", "notes")

COLUMNAS_REPRODUCTORAS = _IDENTIFICACION_AVES + (
    "live_males", "live_females", "birds_placed_males", "birds_placed_females",
    "historical_mortality_males", "historical_mortality_females",
    "historical_culls_males", "historical_culls_females",
    "accumulated_feed_kg", "current_avg_weight_g",
    "accumulated_egg_production", "accumulated_eggs_to_hatchery",
) + _PROCEDENCIA

COLUMNAS_ENGORDE = _IDENTIFICACION_AVES + (
    "live_birds", "age_days", "accumulated_birds_received", "historical_mortality",
    "historical_culls", "partial_outbound_birds", "accumulated_feed_kg", "current_avg_weight_g",
) + _PROCEDENCIA

COLUMNAS_INCUBADORA = (
    "legacy_lot_code", "existing_lot_code", "hatchery_code", "incubator_name", "hatcher_name",
    "source_lot_code", "generation", "egg_origin", "reception_date", "set_date", "expected_hatch_date",
    "eggs_received", "eggs_set", "eggs_transferred", "eggs_discarded",
    "chicks_hatched", "chicks_available",
) + _PROCEDENCIA

COLUMNAS = {
    "grandparent": COLUMNAS_REPRODUCTORAS,
    "breeder": COLUMNAS_REPRODUCTORAS,
    "hatchery": COLUMNAS_INCUBADORA,
    "broiler": COLUMNAS_ENGORDE,
}

#: Maestros propios de la suite v2 (se crean si faltan; códigos únicos del prefijo CUT2).
FASES = (("CRIA", "Cría", 1, True, False), ("PROD", "Producción", 2, False, False),
         ("INCUB", "Incubación", 3, False, False), ("ENGORDE", "Engorde", 4, False, True))
GRANJA = "TEST-GRANJA-01"          # sembrada (seeds.test_seeds.TEST_FARM_CODE)
GALPON = "TEST-GALPON-01"          # sembrado (TEST_HOUSE_NAME), pertenece a GRANJA
PLANTA = "CUT2-INC"
INCUBADORA = "CUT2-INCUBADORA-1"
NACEDORA = "CUT2-NACEDORA-1"
LINEA = "CUT2-LINE"
RAZA = "CUT2-RAZA"
AREA = "CUT2-AREA"

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def corte(dias: int = 40) -> str:
    """Corte civil `dias` atrás a medianoche UTC (sin literales que caduquen)."""
    return days_ago(dias).isoformat() + "T00:00:00+00:00"


def _celda(v: Any) -> Any:
    return "" if v is None else v


def fila(bu: str, **valores: Any) -> list:
    """Fila de `Datos` en el orden del contrato; lo no indicado queda vacío (UNKNOWN)."""
    desconocidas = set(valores) - set(COLUMNAS[bu])
    assert not desconocidas, f"columnas fuera del contrato {bu}: {sorted(desconocidas)}"
    return [_celda(valores.get(c)) for c in COLUMNAS[bu]]


def xlsx_v2(bu: str, filas: Iterable[list], *, company_id: int, template_version: Optional[str] = None,
            cutover_datetime: Optional[str] = None, columnas: Optional[tuple] = None,
            business_unit: Optional[str] = None, con_catalogos: bool = True,
            schema_id: Optional[str] = None) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    instrucciones = wb.active
    instrucciones.title = "Instrucciones"
    instrucciones.append(["Plantilla de prueba SPEC 006"])
    meta = wb.create_sheet("Meta")
    meta.append(["key", "value"])
    version_efectiva = template_version or TEMPLATE_VERSIONS.get(bu, TEMPLATE_VERSION)
    meta.append(["template_version", version_efectiva])
    meta.append(["schema_id", schema_id if schema_id is not None else SCHEMA_IDS.get(bu, "")])
    meta.append(["company", company_id])
    meta.append(["business_unit", business_unit or bu])
    meta.append(["cutover_datetime", cutover_datetime or ""])
    datos = wb.create_sheet("Datos")
    datos.append(list(columnas or COLUMNAS[bu]))
    for f in filas:
        datos.append(list(f))
    if con_catalogos:
        wb.create_sheet("Catalogos").append(["catalog", "code", "description", "status", "parent_code"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_v2(bu: str, filas: Iterable[list], *, delimitador: str = ";", columnas: Optional[tuple] = None,
           bom: bool = True) -> bytes:
    buf = io.StringIO()
    escritor = csv.writer(buf, delimiter=delimitador, lineterminator="\n")
    escritor.writerow(list(columnas or COLUMNAS[bu]))
    for f in filas:
        escritor.writerow(["" if v is None else v for v in f])
    texto = buf.getvalue()
    return ("﻿" + texto if bom else texto).encode("utf-8")


def archivo(contenido: bytes, nombre: str = "corte.xlsx") -> dict:
    mime = "text/csv" if nombre.endswith(".csv") else XLSX_MIME
    return {"file": (nombre, contenido, mime)}


async def crear_batch(client, headers, bu: str, cutover: str, *, template_version: Optional[str] = None) -> int:
    cuerpo = {"business_unit": bu, "cutover_datetime": cutover}
    r = await client.post("/api/v1/cutover-batches", headers=headers, json=cuerpo)
    assert r.status_code == 201, r.text
    return r.json()["id"]


async def subir(client, headers, batch_id: int, contenido: bytes, nombre: str = "corte.xlsx", *,
                template_version: Optional[str] = None):
    url = f"/api/v1/cutover-batches/{batch_id}/upload"
    if template_version:
        url += f"?template_version={template_version}"
    return await client.post(url, headers=headers, files=archivo(contenido, nombre))


async def aprobar(client, headers, headers_aprobador, batch_id: int) -> None:
    r = await client.post(f"/api/v1/cutover-batches/{batch_id}/submit", headers=headers)
    assert r.status_code == 200, r.text
    r = await client.post(f"/api/v1/cutover-batches/{batch_id}/approve", headers=headers_aprobador)
    assert r.status_code == 200, r.text


async def aplicar(client, headers, batch_id: int):
    return await client.post(f"/api/v1/cutover-batches/{batch_id}/apply", headers=headers)


async def items_por_referencia(client, headers, batch_id: int) -> dict[str, int]:
    r = await client.get(f"/api/v1/cutover-batches/{batch_id}/items", headers=headers)
    assert r.status_code == 200, r.text
    return {it["legacy_lot_reference"]: it["lot_id"] for it in r.json()["items"]}


# ── Maestros ─────────────────────────────────────────────────────────────────

async def asegurar_fases() -> dict[str, int]:
    """CRIA/PROD/INCUB/ENGORDE por código (idempotente; una sola fila por código)."""
    from app.masters.models import ProductivePhase

    ids: dict[str, int] = {}
    async with database.async_session() as s:
        for code, name, orden, inicial, final in FASES:
            fase = (await s.execute(select(ProductivePhase).where(ProductivePhase.code == code))).scalars().first()
            if fase is None:
                fase = ProductivePhase(name=name, code=code, order=orden, is_initial=inicial, is_final=final,
                                       is_active=True)
                s.add(fase)
                await s.flush()
            ids[code] = fase.id
        await s.commit()
    return ids


async def asegurar_maestros(company_id: int) -> dict[str, int]:
    """Planta de incubación (+incubadora, +nacedora), línea genética, raza y área de la empresa."""
    from app.masters.models import (Area, BirdTypeEnum, Breed, Farm, GeneticLine, Hatcher, Hatchery,
                                    House, Incubator)

    ids: dict[str, int] = {}
    async with database.async_session() as s:
        granja = (await s.execute(select(Farm).where(Farm.company_id == company_id,
                                                     Farm.code == GRANJA))).scalars().first()
        assert granja is not None, "la granja sembrada TEST-GRANJA-01 debe existir"
        ids["farm_id"] = granja.id
        galpon = (await s.execute(select(House).where(House.farm_id == granja.id,
                                                      House.name == GALPON))).scalars().first()
        assert galpon is not None
        ids["house_id"] = galpon.id

        planta = (await s.execute(select(Hatchery).where(Hatchery.company_id == company_id,
                                                         Hatchery.code == PLANTA))).scalars().first()
        if planta is None:
            planta = Hatchery(company_id=company_id, name="CUT2 Planta", code=PLANTA, is_active=True)
            s.add(planta)
            await s.flush()
        ids["hatchery_id"] = planta.id
        for modelo, nombre, clave in ((Incubator, INCUBADORA, "incubator_id"), (Hatcher, NACEDORA, "hatcher_id")):
            fila_ = (await s.execute(select(modelo).where(modelo.hatchery_id == planta.id,
                                                          modelo.name == nombre))).scalars().first()
            if fila_ is None:
                fila_ = modelo(hatchery_id=planta.id, name=nombre, capacity=100000, is_active=True)
                s.add(fila_)
                await s.flush()
            ids[clave] = fila_.id

        linea = (await s.execute(select(GeneticLine).where(GeneticLine.code == LINEA))).scalars().first()
        if linea is None:
            linea = GeneticLine(company_id=company_id, name="CUT2 Línea", code=LINEA, is_active=True)
            s.add(linea)
            await s.flush()
        ids["genetic_line_id"] = linea.id
        raza = (await s.execute(select(Breed).where(Breed.genetic_line_id == linea.id,
                                                    Breed.name == RAZA))).scalars().first()
        if raza is None:
            raza = Breed(genetic_line_id=linea.id, name=RAZA, bird_type=BirdTypeEnum.BREEDER, is_active=True)
            s.add(raza)
            await s.flush()
        ids["breed_id"] = raza.id
        area = (await s.execute(select(Area).where(Area.company_id == company_id,
                                                   Area.code == AREA))).scalars().first()
        if area is None:
            area = Area(company_id=company_id, name="CUT2 Área", code=AREA, is_active=True)
            s.add(area)
            await s.flush()
        ids["area_id"] = area.id
        await s.commit()
    return ids


async def preparar(seeded_ids: dict) -> dict[str, int]:
    """Fases + maestros + permisos del aprobador (patrón de las suites GA-REQ-061)."""
    from test_ga_req_061_cutover_c3_lifecycle import _otorgar  # noqa: E402

    ids = await asegurar_fases()
    ids.update(await asegurar_maestros(seeded_ids["company_id"]))
    await _otorgar(seeded_ids["role_approver_id"], ("cutover", "approve"), ("cutover", "read"))
    return ids


def headers_aprobador(seeded_ids: dict) -> dict:
    from test_ga_req_061_cutover_c3_lifecycle import _headers_aprobar  # noqa: E402

    return _headers_aprobar(seeded_ids)


# ── Filas canónicas por BU ───────────────────────────────────────────────────

def fila_reproductoras(bu: str, ref: str, *, fase: str = "PROD", **extra: Any) -> list:
    base = dict(legacy_lot_code=ref, farm_code=GRANJA, house_name=GALPON, productive_phase_code=fase,
                real_start_date=iso_days_ago(200), phase_start_date=None,
                live_males=1000, live_females=9000, historical_mortality_males=50, historical_mortality_females=450,
                historical_culls_males=0, historical_culls_females=0, accumulated_feed_kg=120000.5,
                current_avg_weight_g=3200, accumulated_egg_production=180000, accumulated_eggs_to_hatchery=150000,
                source_system="LEGACY", source_reference=f"REF-{ref}")
    base.update(extra)
    return fila(bu, **base)


def fila_engorde(ref: str, **extra: Any) -> list:
    base = dict(legacy_lot_code=ref, farm_code=GRANJA, house_name=GALPON, productive_phase_code="ENGORDE",
                real_start_date=iso_days_ago(150), live_birds=10000, age_days=20, accumulated_birds_received=10500,
                historical_mortality=300, historical_culls=200, partial_outbound_birds=0,
                accumulated_feed_kg=15000, current_avg_weight_g=900, source_system="LEGACY")
    base.update(extra)
    return fila("broiler", **base)


def fila_incubadora(ref: str, **extra: Any) -> list:
    base = dict(legacy_lot_code=ref, hatchery_code=PLANTA, incubator_name=INCUBADORA, hatcher_name=NACEDORA,
                generation="breeder", egg_origin="internal", reception_date=iso_days_ago(120), set_date=iso_days_ago(118),
                expected_hatch_date=iso_days_ago(97), eggs_received=10000, eggs_set=9000,
                eggs_transferred=0, eggs_discarded=100, chicks_hatched=0, chicks_available=0,
                source_system="LEGACY")
    base.update(extra)
    return fila("hatchery", **base)
