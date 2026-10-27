"""SPEC 011 CUTOVER-DATA-01 — Meta refleja el corte (display) pero no es autoridad.

El endpoint `GET /cutover-templates/{bu}?cutover_datetime=` escribe la fecha de corte en
`Meta.cutover_datetime` para que el usuario vea el corte de su batch; en cambio el parser
**ignora** ese valor: la autoridad es `batch.cutover_datetime` (lo usan apply y reconciliación).

AC: 2 funciones de test definidas; el módulo importa sin error.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from cutover_v2_fixtures import fila_incubadora, preparar, xlsx_v2  # noqa: E402

from app.cutover.contracts import contrato  # noqa: E402
from app.cutover.parser import parse_cutover_file  # noqa: E402


def _meta(contenido_bytes: bytes) -> dict:
    wb = load_workbook(io.BytesIO(contenido_bytes))
    return {str(k): v for k, v in wb["Meta"].iter_rows(values_only=True) if k is not None}


@pytest.mark.asyncio
async def test_data01_meta_refleja_corte_sin_batch(auth_headers, http_client, seeded_ids):
    """Sin `batch_id`, `?cutover_datetime=` se refleja en `Meta` (display para el usuario)."""
    await preparar(seeded_ids)
    r = await http_client.get("/api/v1/cutover-templates/hatchery?cutover_datetime=2026-10-01T08:00:00%2B00:00",
                              headers=auth_headers)
    assert r.status_code == 200, r.text
    meta = _meta(r.content)
    assert str(meta.get("cutover_datetime", "")).startswith("2026-10-01T08:00:00")  # Meta refleja el corte (display)


def test_data01_parser_ignora_cutover_de_meta():
    """El parser NO toma la fecha de corte de Meta: la autoridad es el batch."""
    contenido = xlsx_v2("hatchery", [fila_incubadora("S011-DATA01")], company_id=4,
                        cutover_datetime="2099-12-31T00:00:00+00:00")
    dto = parse_cutover_file(contenido, "corte.xlsx", contrato=contrato("hatchery"))
    assert "cutover_datetime" not in dto  # el DTO del parser no expone cutover_datetime de Meta
    assert dto["business_unit"] == "hatchery" and dto["template_version"] == "v3"
