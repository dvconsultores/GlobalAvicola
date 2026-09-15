"""GA-REQ-061 · T14 · C9 — plantillas Excel por BU (descarga).

`GA_CUTOVER_EXCEL_STAGING_DESIGN §2`: plantilla por BU, versionada
(`cutover_template_{bu}_v1.xlsx`), con hojas `Instrucciones`, `Meta` (con
`template_version`, `company`, `business_unit`, `cutover_datetime`) y `Datos`.
La semántica UNKNOWN viaja escrita en la propia plantilla: **celda vacía =
UNKNOWN · `0` explícito = cero conocido · `N/A` = NOT_APPLICABLE** (AC15/AC77).

Gate: el mismo de crear batch (`cutover:create`, matriz de seguridad §5).
"""
from __future__ import annotations

import io

from fastapi import APIRouter, Depends, HTTPException, Response, status

from ..dependencies import require_permission
from ..transaction import RutaTransaccional
from .parser import SUPPORTED_TEMPLATE_VERSIONS

router = APIRouter(route_class=RutaTransaccional, prefix="/cutover-templates", tags=["Cutover"])

UNIDADES = ("grandparent", "breeder", "hatchery", "broiler")

COLUMNAS_DATOS = (
    "legacy_lot_code", "real_start_date", "live_males", "live_females",
    "historical_mortality_males", "historical_mortality_females", "farm_code", "notes",
)

INSTRUCCIONES = (
    "PLANTILLA DE CORTE OPERACIONAL (Cargas Iniciales) — cómo se llena",
    "",
    "Semántica del dato histórico (se conserva tal cual, jamás se fabrica):",
    "  · celda vacía  = UNKNOWN (sin dato en el sistema de origen)",
    "  · 0 explícito  = cero CONOCIDO",
    "  · N/A          = NOT_APPLICABLE para la unidad",
    "",
    "Columnas de la hoja Datos:",
    "  legacy_lot_code  código del lote en el sistema anterior (obligatorio)",
    "  real_start_date  fecha real de inicio del lote (ISO AAAA-MM-DD)",
    "  live_males       aves vivas machos al momento del corte",
    "  live_females     aves vivas hembras al momento del corte",
    "  historical_mortality_males / _females  mortalidad acumulada previa",
    "  farm_code        código de granja existente (opcional)",
    "  notes            notas libres",
    "",
    "No use fórmulas ni formatos como dato: el importador lee valores.",
)


def construir_plantilla(business_unit: str, company_id: int, template_version: str = "v1") -> bytes:
    """Genera la plantilla v1 de la BU (mismas columnas que el parser v1 soporta)."""
    from openpyxl import Workbook

    if template_version not in SUPPORTED_TEMPLATE_VERSIONS:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                            detail=f"TEMPLATE_VERSION_UNSUPPORTED: {template_version}")

    wb = Workbook()
    instrucciones = wb.active
    instrucciones.title = "Instrucciones"
    for linea in INSTRUCCIONES:
        instrucciones.append([linea])

    meta = wb.create_sheet("Meta")
    meta.append(["key", "value"])
    meta.append(["template_version", template_version])
    meta.append(["company", company_id])
    meta.append(["business_unit", business_unit])
    meta.append(["cutover_datetime", ""])  # lo fija el batch al subir

    datos = wb.create_sheet("Datos")
    datos.append(list(COLUMNAS_DATOS))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/{business_unit}")
async def descargar_plantilla(
    business_unit: str,
    current_user: dict = Depends(require_permission("cutover", "create")),
):
    """`GET /cutover-templates/{bu}` — descarga la plantilla versionada de la BU."""
    if business_unit not in UNIDADES:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                            detail=f"BUSINESS_UNIT_INVALID: '{business_unit}' no es una unidad de negocio")
    empresa = current_user.get("effective_company_id") or current_user.get("company_id")
    if not empresa:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN,
                            detail="Se requiere una empresa efectiva para descargar la plantilla.")
    contenido = construir_plantilla(business_unit, int(empresa))
    return Response(
        content=contenido,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="cutover_template_{business_unit}_v1.xlsx"'},
    )
