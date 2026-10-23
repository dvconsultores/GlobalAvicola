"""SPEC 006 · Cutover v2 — parser XLSX + CSV hacia un único DTO normalizado.

El archivo es **dato, jamás código**: fórmulas y macros se rechazan sin evaluarse. Ambos
formatos producen exactamente la misma estructura (`XLSX_RESULT == CSV_RESULT`):

    {"source_type": "XLSX"|"CSV", "template_version", "business_unit", "company", "schema_id",
     "rows": [{"row_number", "raw": {col: str|None}, "normalized": {campo: {"value", "status"}},
               "errors": [{row_number, column, field, error_code, message, received_value}]}]}

Semántica de celda (contrato): **vacío = UNKNOWN · `0` = KNOWN cero · `N/A` = NOT_APPLICABLE**.
`N/A` donde el dato es obligatorio ⇒ `NOT_APPLICABLE_ON_REQUIRED` (nunca un cero fabricado).
La versión `v1` está retirada: se rechaza con `TEMPLATE_VERSION_UNSUPPORTED` y mensaje claro.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime
from typing import Any, Optional

from .contracts import (KNOWN, NOT_APPLICABLE, RETIRED_TEMPLATE_VERSIONS, SUPPORTED_TEMPLATE_VERSIONS,
                        TEMPLATE_VERSION, UNKNOWN, FieldSpec, TemplateContract)

#: Límites de archivo (mandato §W).
LIMITE_BYTES = 5 * 1024 * 1024
LIMITE_FILAS = 5000
LIMITE_COLUMNAS = 60

_NA = {"N/A", "NA", "N.A.", "N.A"}
_ENTERO = re.compile(r"^-?\d+$")
_ENTERO_CON_CERO = re.compile(r"^-?\d+\.0+$")
_DECIMAL = re.compile(r"^-?\d+(\.\d+)?$")
_FECHA = re.compile(r"^\d{4}-\d{2}-\d{2}$")


class CutoverParseError(Exception):
    """Error de archivo/plantilla — viaja como `error_code` estructurado (422)."""

    def __init__(self, code: str, message: str):
        super().__init__(f"{code}: {message}")
        self.code = code
        self.message = message


def error_de_fila(row: int, column: str, field: str, code: str, message: str, received: Any) -> dict:
    return {
        "row_number": row, "column": column, "field": field, "error_code": code, "message": message,
        "received_value": None if received is None else str(received),
    }


def _texto(v: Any) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _celda_vacia(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


# ── Lectura ───────────────────────────────────────────────────────────────────

def _extension(filename: str) -> str:
    nombre = (filename or "").lower()
    return nombre.rsplit(".", 1)[-1] if "." in nombre else ""


def _leer_xlsx(content: bytes) -> tuple[dict, list[str], list[tuple[int, list]]]:
    """Devuelve (meta, cabecera, filas[(n, celdas)]). Rechaza macros y fórmulas sin evaluarlas."""
    if not content.startswith(b"PK"):
        raise CutoverParseError("MIME_MISMATCH", "El archivo no es un libro XLSX válido (la extensión no coincide con el contenido).")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            nombres = set(z.namelist())
    except zipfile.BadZipFile as exc:
        raise CutoverParseError("INVALID_FILE", "XLSX corrupto: no se pudo abrir el contenedor.") from exc
    if any(n.lower().endswith("vbaproject.bin") for n in nombres):
        raise CutoverParseError("MACRO_NOT_ALLOWED", "El archivo contiene macros; la plantilla es solo datos.")
    try:
        from openpyxl import load_workbook

        libro = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:  # corrupto
        raise CutoverParseError("INVALID_FILE", f"No se pudo leer el archivo XLSX: {exc}") from exc

    for hoja in ("Meta", "Datos"):
        if hoja not in libro.sheetnames:
            raise CutoverParseError("SHEET_MISSING", f"El archivo no trae la hoja '{hoja}'. Descargue la plantilla v2 de la unidad.")

    meta: dict[str, Any] = {}
    for fila in libro["Meta"].iter_rows(values_only=True):
        if fila and _texto(fila[0]):
            meta[str(fila[0]).strip()] = (fila[1] if len(fila) > 1 else None)

    filas_crudas = list(libro["Datos"].iter_rows(values_only=False))
    if not filas_crudas:
        raise CutoverParseError("HEADER_MISSING", "La hoja 'Datos' está vacía: falta la fila de cabecera.")
    cabecera_celdas = filas_crudas[0]
    cabecera = [str(c.value).strip() if c.value is not None else "" for c in cabecera_celdas]
    while cabecera and cabecera[-1] == "":
        cabecera.pop()
    filas: list[tuple[int, list]] = []
    for indice, celdas in enumerate(filas_crudas[1:], start=2):
        valores = []
        for c in celdas[:len(cabecera)]:
            v = c.value
            if c.data_type == "f" or (isinstance(v, str) and v.startswith("=")):
                raise CutoverParseError(
                    "FORMULA_NOT_ALLOWED",
                    f"Fórmula en la fila {indice}, columna {c.column_letter}: la plantilla acepta valores, no fórmulas.")
            valores.append(v)
        filas.append((indice, valores))
    libro.close()
    return meta, cabecera, filas


def _leer_csv(content: bytes) -> tuple[list[str], list[tuple[int, list]]]:
    try:
        texto = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise CutoverParseError("INVALID_FILE", "CSV corrupto o con codificación no soportada (se espera UTF-8).") from exc
    if "\x00" in texto:
        raise CutoverParseError("INVALID_FILE", "CSV corrupto: contiene bytes nulos.")
    lineas = texto.splitlines()
    primera = next((l for l in lineas if l.strip()), None)
    if primera is None:
        raise CutoverParseError("HEADER_MISSING", "El CSV está vacío: falta la fila de cabecera.")
    delimitador = ";" if primera.count(";") >= primera.count(",") else ","
    if "\t" in primera and primera.count("\t") > max(primera.count(";"), primera.count(",")):
        delimitador = "\t"
    lector = csv.reader(io.StringIO(texto), delimiter=delimitador)
    cabecera = [c.strip() for c in next(lector)]
    while cabecera and cabecera[-1] == "":
        cabecera.pop()
    filas: list[tuple[int, list]] = []
    for indice, valores in enumerate(lector, start=2):
        filas.append((indice, [v if v is not None else "" for v in valores[:len(cabecera)]]))
    return cabecera, filas


# ── Cabecera ──────────────────────────────────────────────────────────────────

def _validar_cabecera(cabecera: list[str], contrato: TemplateContract) -> dict[str, int]:
    if len(cabecera) > LIMITE_COLUMNAS:
        raise CutoverParseError("TOO_MANY_COLUMNS", f"La plantilla trae {len(cabecera)} columnas; el máximo es {LIMITE_COLUMNAS}.")
    vistas: dict[str, int] = {}
    duplicadas: list[str] = []
    for i, nombre in enumerate(cabecera):
        if not nombre:
            continue
        if nombre in vistas:
            duplicadas.append(nombre)
        else:
            vistas[nombre] = i
    if duplicadas:
        raise CutoverParseError("HEADER_DUPLICATE", f"Columnas repetidas en la cabecera: {', '.join(sorted(set(duplicadas)))}.")
    desconocidas = sorted(set(vistas) - set(contrato.columns))
    if desconocidas:
        raise CutoverParseError("HEADER_UNKNOWN",
                                f"Columnas que no pertenecen al contrato {contrato.schema_id}: {', '.join(desconocidas)}.")
    faltantes = [c for c in contrato.columns if c not in vistas]
    if faltantes:
        raise CutoverParseError("HEADER_MISSING",
                                f"Faltan columnas del contrato {contrato.schema_id}: {', '.join(faltantes)}. Descargue la plantilla v2.")
    return vistas


# ── Normalización por campo ───────────────────────────────────────────────────

def _normalizar_valor(spec: FieldSpec, v: Any) -> tuple[Any, Optional[str]]:
    """Devuelve (valor, código de error). `v` no está vacío ni es N/A."""
    if spec.kind == "text":
        s = str(v).strip() if not isinstance(v, str) else v.strip()
        if isinstance(v, float) and v.is_integer():
            s = str(int(v))
        if spec.max_len and len(s) > spec.max_len:
            return None, "INVALID_TEXT_LENGTH"
        return s, None
    if spec.kind == "enum":
        s = str(v).strip().lower()
        return (s, None) if s in spec.values else (None, "INVALID_ENUM")
    if spec.kind == "int":
        if isinstance(v, bool):
            return None, "INVALID_INTEGER"
        if isinstance(v, int):
            n = v
        elif isinstance(v, float):
            if not v.is_integer():
                return None, "INVALID_INTEGER"
            n = int(v)
        else:
            s = str(v).strip()
            if _ENTERO.match(s):
                n = int(s)
            elif _ENTERO_CON_CERO.match(s):
                n = int(float(s))
            else:
                return None, "INVALID_INTEGER"
        if n < 0:
            return None, "NEGATIVE_VALUE"
        return n, None
    if spec.kind == "decimal":
        if isinstance(v, bool):
            return None, "INVALID_NUMBER"
        if isinstance(v, (int, float)):
            x = float(v)
        else:
            s = str(v).strip()
            if not _DECIMAL.match(s):
                return None, "INVALID_NUMBER"
            x = float(s)
        if x < 0:
            return None, "NEGATIVE_VALUE"
        if spec.positive and x <= 0:
            return None, "INVALID_NUMBER"
        return x, None
    if spec.kind == "date":
        if isinstance(v, datetime):
            return v.date().isoformat(), None
        if isinstance(v, date):
            return v.isoformat(), None
        s = str(v).strip()
        if not _FECHA.match(s):
            return None, "INVALID_DATE"
        try:
            return date.fromisoformat(s).isoformat(), None
        except ValueError:
            return None, "INVALID_DATE"
    return None, "INVALID_FILE"


_MENSAJES = {
    "REQUIRED_FIELD_MISSING": "Campo obligatorio sin valor.",
    "NOT_APPLICABLE_ON_REQUIRED": "N/A no es válido en un campo obligatorio: indique el valor (0 si es cero real).",
    "NOT_APPLICABLE_NOT_ALLOWED": "N/A no es válido para este campo.",
    "INVALID_INTEGER": "Número entero ≥ 0 esperado (sin separador de miles ni decimales).",
    "INVALID_NUMBER": "Número decimal ≥ 0 esperado, con punto decimal.",
    "NEGATIVE_VALUE": "El valor no puede ser negativo.",
    "INVALID_DATE": "Fecha inválida (se espera AAAA-MM-DD).",
    "INVALID_ENUM": "Valor fuera de la lista permitida.",
    "INVALID_TEXT_LENGTH": "Texto demasiado largo.",
}


def normalizar_fila(contrato: TemplateContract, numero: int, columnas: dict[str, int], valores: list) -> dict:
    crudos: dict[str, Optional[str]] = {}
    normalizado: dict[str, dict] = {}
    errores: list[dict] = []
    for spec in contrato.fields:
        i = columnas[spec.name]
        v = valores[i] if i < len(valores) else None
        crudos[spec.name] = None if _celda_vacia(v) else str(v)
        if _celda_vacia(v):
            if spec.required:
                errores.append(error_de_fila(numero, spec.name, spec.name, "REQUIRED_FIELD_MISSING",
                                             _MENSAJES["REQUIRED_FIELD_MISSING"], None))
            else:
                normalizado[spec.normalized] = {"value": None, "status": UNKNOWN}
            continue
        if isinstance(v, str) and v.strip().upper() in _NA:
            if spec.required:
                errores.append(error_de_fila(numero, spec.name, spec.name, "NOT_APPLICABLE_ON_REQUIRED",
                                             _MENSAJES["NOT_APPLICABLE_ON_REQUIRED"], v))
            elif not spec.allow_na:
                errores.append(error_de_fila(numero, spec.name, spec.name, "NOT_APPLICABLE_NOT_ALLOWED",
                                             _MENSAJES["NOT_APPLICABLE_NOT_ALLOWED"], v))
            else:
                normalizado[spec.normalized] = {"value": None, "status": NOT_APPLICABLE}
            continue
        valor, codigo = _normalizar_valor(spec, v)
        if codigo:
            mensaje = _MENSAJES.get(codigo, codigo)
            if codigo == "INVALID_ENUM":
                mensaje = f"Valor fuera de la lista permitida ({', '.join(spec.values)})."
            errores.append(error_de_fila(numero, spec.name, spec.name, codigo, mensaje, v))
            continue
        normalizado[spec.normalized] = {"value": valor, "status": KNOWN}
    return {"row_number": numero, "raw": crudos, "normalized": normalizado, "errors": errores}


# ── Punto de entrada ──────────────────────────────────────────────────────────

def parse_cutover_file(content: bytes, filename: str, *, contrato: TemplateContract,
                       template_version: Optional[str] = None) -> dict:
    """Parsea XLSX o CSV contra el contrato de la BU del batch.

    `template_version` solo se usa para CSV (sin hoja Meta): la fija el endpoint (`?template_version=`, por
    defecto `v2`). Para XLSX la versión viaja en `Meta`.
    """
    if not content:
        raise CutoverParseError("INVALID_FILE", "El archivo está vacío.")
    if len(content) > LIMITE_BYTES:
        raise CutoverParseError("FILE_TOO_LARGE", f"El archivo supera el máximo de {LIMITE_BYTES // (1024 * 1024)} MB.")
    ext = _extension(filename)
    if ext == "xlsm":
        raise CutoverParseError("MACRO_NOT_ALLOWED", "Los libros con macros (.xlsm) no se aceptan; use .xlsx o .csv.")
    if ext not in ("xlsx", "csv"):
        raise CutoverParseError("UNSUPPORTED_EXTENSION", f"Extensión '.{ext}' no soportada: use .xlsx o .csv.")

    if ext == "xlsx":
        meta, cabecera, filas = _leer_xlsx(content)
        version = _texto(meta.get("template_version"))
        business_unit = _texto(meta.get("business_unit"))
        schema_id = _texto(meta.get("schema_id"))
        company = meta.get("company")
        source_type = "XLSX"
    else:
        if content.startswith(b"PK"):
            raise CutoverParseError("MIME_MISMATCH", "El archivo tiene extensión .csv pero su contenido es un libro XLSX.")
        cabecera, filas = _leer_csv(content)
        version = template_version or TEMPLATE_VERSION
        business_unit = contrato.business_unit
        schema_id = contrato.schema_id
        company = None
        source_type = "CSV"

    if version in RETIRED_TEMPLATE_VERSIONS:
        raise CutoverParseError(
            "TEMPLATE_VERSION_UNSUPPORTED",
            f"La plantilla {version} está retirada: descargue la plantilla {TEMPLATE_VERSION} de {contrato.business_unit} y vuelva a cargar.")
    if version not in SUPPORTED_TEMPLATE_VERSIONS:
        raise CutoverParseError("TEMPLATE_VERSION_UNSUPPORTED",
                                f"Versión de plantilla no soportada: {version!r} (soportadas: {sorted(SUPPORTED_TEMPLATE_VERSIONS)}).")
    if business_unit == contrato.business_unit and version != contrato.version:
        raise CutoverParseError(
            "TEMPLATE_VERSION_UNSUPPORTED",
            f"La plantilla de {contrato.business_unit} vigente es {contrato.version}: descargue la plantilla "
            f"{contrato.version} de {contrato.business_unit} y vuelva a cargar.")
    if not business_unit:
        raise CutoverParseError("BUSINESS_UNIT_MISMATCH", "La hoja 'Meta' no declara 'business_unit'.")
    if business_unit != contrato.business_unit or (schema_id and schema_id != contrato.schema_id):
        raise CutoverParseError(
            "BUSINESS_UNIT_MISMATCH",
            f"La plantilla es de {business_unit} ({schema_id or 'sin schema_id'}) y el batch es de {contrato.business_unit}.")

    columnas = _validar_cabecera(cabecera, contrato)

    resultado: list[dict] = []
    for numero, valores in filas:
        if all(_celda_vacia(c) for c in valores):
            continue
        if len(resultado) >= LIMITE_FILAS:
            raise CutoverParseError("TOO_MANY_ROWS", f"La plantilla supera el máximo de {LIMITE_FILAS} filas de datos.")
        resultado.append(normalizar_fila(contrato, numero, columnas, valores))
    if not resultado:
        raise CutoverParseError("NO_DATA_ROWS", "La plantilla no trae filas de datos: complete la hoja 'Datos' antes de subirla.")

    empresa: Optional[int] = None
    if company not in (None, ""):
        try:
            empresa = int(float(str(company).strip()))
        except ValueError:
            raise CutoverParseError("COMPANY_MISMATCH", f"Meta.company inválido: {company!r}.")

    return {"source_type": source_type, "template_version": version, "business_unit": business_unit,
            "schema_id": schema_id, "company": empresa, "rows": resultado}

