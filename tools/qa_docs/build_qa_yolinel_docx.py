"""Convierte los documentos Markdown de la revisión QA (Yolinel Cortez 2026-10-17) a DOCX.

Uso: backend/.venv/bin/python tools/qa_docs/build_qa_yolinel_docx.py
Genera los DOCX en docs/ y reporta "OK <docx> párrafos=<n> tablas=<m>" por documento.
Si un .md no existe imprime "OMITIDO <ruta>" y continúa (rc 0).
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]

DOCUMENTOS = [
    (
        "audit/qa-yolinel-2026-10-17/09_RESPONSE_TO_QA.md",
        "docs/RESPUESTA_REVISION_GENERAL_QA_YOLINEL_2026_10_17.docx",
        "Respuesta a la Revisión General de Sistema — QA Yolinel Cortez (17-10-2026)",
    ),
    (
        "audit/qa-yolinel-2026-10-17/OPERATION_STATE_MACHINE.md",
        "docs/FLUJO_DE_ESTADOS_GLOBAL_AVICOLA.docx",
        "Flujo de estados de los eventos operativos — Global Avícola",
    ),
    (
        "audit/qa-yolinel-2026-10-17/PLAN_QA_REPRODUCTORAS_INCUBADORA.md",
        "docs/PLAN_QA_REPRODUCTORAS_INCUBADORA.docx",
        "Plan de QA — Reproductoras e Incubadora",
    ),
]

FECHA = "2026-10-18"
MONO = "Consolas"


def _inline(par, texto, cursiva=False):
    """**negrita**, `código` (Consolas) y *cursiva* → runs; limpia enlaces y comentarios HTML."""
    texto = re.sub(r"<!--.*?-->", "", texto)
    texto = re.sub(r"!?\[([^\]]*)\]\([^)]*\)", r"\1", texto)
    for parte in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)", texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            r = par.add_run(parte[2:-2])
            r.bold = True
            r.italic = cursiva
        elif parte.startswith("`") and parte.endswith("`") and len(parte) > 2:
            r = par.add_run(parte[1:-1])
            r.font.name = MONO
            r.font.size = Pt(9)
            r.italic = cursiva
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            r = par.add_run(parte[1:-1])
            r.italic = True
        else:
            r = par.add_run(parte)
            r.italic = cursiva
    return par


def _celdas(fila):
    return [c.strip() for c in fila.strip().strip("|").split("|")]


def _tabla(doc, filas):
    # algunas tablas vienen escritas en una sola línea con separadores literales "\n"
    filas = [sub for f in filas for sub in f.split("\\n") if sub.strip()]
    cuerpo = [f for f in filas if not re.match(r"^\s*\|?[\s:\-|]+$", f)]
    if not cuerpo:
        return
    matriz = [_celdas(f) for f in cuerpo]
    ncols = max(len(f) for f in matriz)
    t = doc.add_table(rows=len(matriz), cols=ncols)
    t.style = "Light Grid Accent 1"
    for i, fila in enumerate(matriz):
        fila = fila + [""] * (ncols - len(fila))  # rellena celdas desiguales
        for j, val in enumerate(fila):
            p = t.cell(i, j).paragraphs[0]
            _inline(p, val)
            for r in p.runs:
                r.font.size = Pt(9)
                if i == 0:
                    r.bold = True
    doc.add_paragraph()


def md_a_docx(md_path, docx_path, titulo):
    """Convierte md_path (Markdown) en docx_path (Word). Devuelve (párrafos, tablas)."""
    md_path = Path(md_path)
    docx_path = Path(docx_path)
    if not md_path.exists():
        print(f"OMITIDO {md_path}")
        return None

    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)

    relativo = md_path.resolve().relative_to(ROOT).as_posix()

    # Portada
    doc.add_heading(titulo, 0)
    portada = doc.add_paragraph()
    _inline(portada, f"Global Avícola · {FECHA} · Generado desde {relativo}")
    for r in portada.runs:
        r.font.size = Pt(9.5)
        r.italic = True

    lineas = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lineas):
        l = lineas[i]
        desnudo = l.strip()

        # Bloques de código ```
        if desnudo.startswith("```"):
            i += 1
            buffer = []
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                buffer.append(lineas[i])
                i += 1
            i += 1  # consume el cierre
            p = doc.add_paragraph()
            r = p.add_run("\n".join(buffer))
            r.font.name = MONO
            r.font.size = Pt(9)
            continue

        # Tablas
        if desnudo.startswith("|"):
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                bloque.append(lineas[i])
                i += 1
            _tabla(doc, bloque)
            continue

        # Encabezados # .. ####
        m = re.match(r"^(#{1,4})\s+(.*)$", desnudo)
        if m:
            doc.add_heading(re.sub(r"[*`]", "", m.group(2).strip()), level=len(m.group(1)))
            i += 1
            continue

        # Cita
        if desnudo.startswith(">"):
            p = doc.add_paragraph()
            _inline(p, desnudo.lstrip(">").strip(), cursiva=True)
            i += 1
            continue

        # Separador horizontal
        if re.match(r"^-{3,}$", desnudo):
            doc.add_paragraph()
            i += 1
            continue

        # Listas
        m = re.match(r"^[-*]\s+(.*)$", desnudo)
        if m:
            _inline(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue
        m = re.match(r"^\d+\.\s+(.*)$", desnudo)
        if m:
            _inline(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue

        # Línea en blanco
        if not desnudo:
            i += 1
            continue

        # Párrafo normal
        _inline(doc.add_paragraph(), desnudo)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    return len(doc.paragraphs), len(doc.tables)


def main():
    fallos = 0
    for md, docx, titulo in DOCUMENTOS:
        resultado = md_a_docx(ROOT / md, ROOT / docx, titulo)
        if resultado is None:
            continue
        parrafos, tablas = resultado
        print(f"OK {docx} párrafos={parrafos} tablas={tablas}")
    return 0 if fallos == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
